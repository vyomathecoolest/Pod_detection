import cv2
import pytesseract
import re
import os
from openpyxl import load_workbook

# Folder containing images
folder_path = r"C:\Users\vyoma\OneDrive\Documents\Desktop\internship\classified_by_lsp\resized_gati"

# Fixed bounding box coordinates
x, y, w, h = 840, 300, 340, 100

# OCR config
custom_config = r'--oem 3 --psm 6'

# Path to your existing Excel file
excel_file_path = r"C:\Users\vyoma\OneDrive\Documents\Desktop\internship\blur_classified\blur_results.xlsx"

# Load the workbook and active worksheet
wb = load_workbook(excel_file_path)
ws = wb.active

# Check if header "Predicted Number" exists; if not, add it in the next empty column in row 1
headers = [cell.value for cell in ws[1]]
if "Predicted Number" not in headers:
    ws.cell(row=1, column=len(headers)+1, value="Predicted Number")
predicted_col = headers.index("Predicted Number") + 1 if "Predicted Number" in headers else len(headers) + 1

def process_image(image_path):
    image = cv2.imread(image_path)
    if image is None:
        print(f"Error: Could not load image {image_path}")
        return []

    cropped = image[y:y+h, x:x+w]
    gray = cv2.cvtColor(cropped, cv2.COLOR_BGR2GRAY)
    resized = cv2.resize(gray, None, fx=2, fy=2, interpolation=cv2.INTER_LINEAR)
    blurred = cv2.GaussianBlur(resized, (3, 3), 0)

    text = pytesseract.image_to_string(blurred, config=custom_config)
    nine_digit_numbers = re.findall(r'\b\d{9}\b', text)
    return nine_digit_numbers

# Create a dict from existing filenames to row numbers for quick lookup
filename_to_row = {}
for row in range(2, ws.max_row + 1):
    cell_value = ws.cell(row=row, column=1).value
    if cell_value:
        filename_to_row[cell_value] = row

# Process each image in folder and update predicted number column
for filename in os.listdir(folder_path):
    if filename.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.tiff')):
        image_path = os.path.join(folder_path, filename)
        numbers = process_image(image_path)
        numbers_str = ', '.join(numbers) if numbers else ''

        # Check if filename exists in Excel, else append new row
        if filename in filename_to_row:
            row_to_write = filename_to_row[filename]
        else:
            row_to_write = ws.max_row + 1
            ws.cell(row=row_to_write, column=1, value=filename)

        # Write predicted numbers in the "Predicted Number" column
        ws.cell(row=row_to_write, column=predicted_col, value=numbers_str)
        print(f"Updated {filename} with predicted numbers: {numbers_str}")

# Save workbook
wb.save(excel_file_path)
print(f"Excel file updated and saved at {excel_file_path}")
